import logging
import os
import re
import shutil
import tempfile
from copy import copy
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.table import Table
from openpyxl.worksheet.worksheet import Worksheet

from .config import CATALOGUE_BACKUP_DIR_NAME, CATALOGUE_SHEET_NAME
from .errors import CatalogueError, InputError
from .schema import DaliaInterchangeFormat
from .service import normalize_title, read_snapshot

logger = logging.getLogger(__name__)

DIF_FIELDS = list(DaliaInterchangeFormat.model_fields)
CATALOGUE_COLUMNS = ["No", *DIF_FIELDS, "Publisher"]

# Excel refuses cells longer than this.
MAX_CELL_LENGTH = 32767

_ORCID = re.compile(r":\s*\{[^}]*\}")
_PARSED_FORMATS = {".json": "json", ".tsv": "tsv"}


@dataclass(frozen=True)
class ParsedRecord:
    """A parsed record together with where it came from, for messages."""

    source: Path
    number: int
    metadata: DaliaInterchangeFormat

    @property
    def label(self) -> str:
        return f"{self.source.name} #{self.number}"


@dataclass(frozen=True)
class Duplicate:
    """A record that is already present, and what it matched."""

    record: ParsedRecord
    found_at: str
    matched_on: str


@dataclass
class MergePlan:
    """What a merge will write, decided before the catalogue is touched."""

    to_write: list[ParsedRecord] = field(default_factory=list)
    skipped: list[Duplicate] = field(default_factory=list)
    forced: list[Duplicate] = field(default_factory=list)


@dataclass(frozen=True)
class MergeReport:
    """The rows a merge wrote."""

    first_row: int
    last_row: int

    @property
    def count(self) -> int:
        return self.last_row - self.first_row + 1


def read_parsed(path: Path) -> list[ParsedRecord]:
    """The records of a `parse` output file (.json or .tsv)."""
    format = _PARSED_FORMATS.get(path.suffix.lower())
    if format is None:
        raise InputError(
            f"Unsupported input '{path.suffix}'. Use one of: {', '.join(_PARSED_FORMATS)}."
        )
    records = read_snapshot(path, format=format)
    if not records:
        raise InputError(f"{path} contains no records.")
    return [
        ParsedRecord(source=path, number=number, metadata=metadata)
        for number, metadata in enumerate(records, start=1)
    ]


def backup_catalogue(path: Path) -> Path:
    """Copy the catalogue into a timestamped file in `.backup/` next to it."""
    backup_dir = path.parent / CATALOGUE_BACKUP_DIR_NAME
    backup_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    target = backup_dir / f"{path.stem}_{stamp}{path.suffix}"
    counter = 1
    while target.exists():
        counter += 1
        target = backup_dir / f"{path.stem}_{stamp}_{counter}{path.suffix}"
    shutil.copy2(path, target)
    return target


def normalize_link(value: str) -> str:
    """A link without scheme, `www.`, trailing slash and case."""
    link = re.sub(r"^https?://", "", value.strip().casefold())
    return link.removeprefix("www.").rstrip("/")


def normalize_authors(value: str) -> tuple[str, ...]:
    """The author names of a `*`-separated list, without ORCIDs, in a stable order."""
    names = (re.sub(r"\s+", " ", name).strip().casefold() for name in _ORCID.sub("", value).split("*"))
    return tuple(sorted(name for name in names if name))


class DuplicateIndex:
    """
    Recognises records already in the catalogue: a record is a duplicate when
    its Link matches, or when both its Title and Authors match.
    """

    def __init__(self) -> None:
        self._links: dict[str, str] = {}
        self._works: dict[tuple[str, tuple[str, ...]], str] = {}

    def add(self, authors: str, link: str, title: str, found_at: str) -> None:
        if link_key := normalize_link(link):
            self._links.setdefault(link_key, found_at)
        title_key, authors_key = normalize_title(title), normalize_authors(authors)
        if title_key and authors_key:
            self._works.setdefault((title_key, authors_key), found_at)

    def find(self, record: ParsedRecord) -> Duplicate | None:
        metadata = record.metadata
        if found_at := self._links.get(normalize_link(metadata.Link)):
            return Duplicate(record, found_at, "Link")
        work = (normalize_title(metadata.Title), normalize_authors(metadata.Authors))
        if found_at := self._works.get(work):
            return Duplicate(record, found_at, "Title + Authors")
        return None


def cell_value(value: str | None) -> str | None:
    """A field value as it is stored in a cell: empty becomes None, illegal characters go."""
    if value is None or value == "":
        return None
    cleaned = ILLEGAL_CHARACTERS_RE.sub("", value)
    if len(cleaned) > MAX_CELL_LENGTH:
        raise CatalogueError(
            f"Value of {len(cleaned)} characters exceeds Excel's cell limit of {MAX_CELL_LENGTH}."
        )
    return cleaned


def row_values(no: int, metadata: DaliaInterchangeFormat) -> list:
    """The cells of one catalogue row, in `CATALOGUE_COLUMNS` order."""
    data = metadata.model_dump()
    return [no, *(cell_value(data[name]) for name in DIF_FIELDS), None]


class CatalogueMerger:
    """
    Appends parsed records to the OER catalogue sheet.

    `plan` decides what to write without touching anything; `write` appends the
    rows, saves to a temporary file, validates it and only then replaces the
    catalogue.
    """

    def __init__(self, path: Path):
        self._path = path
        self._workbook = _open_workbook(path)
        self._sheet = _catalogue_sheet(self._workbook, path)
        _check_header(self._sheet, path)
        self._table = _catalogue_table(self._sheet, path)
        self._last_row = range_boundaries(self._table.ref)[3]
        self._snapshot = _values(self._sheet, 2, self._last_row)
        self._next_no = 1 + max(
            (row[0] for row in self._snapshot if isinstance(row[0], int | float)),
            default=0,
        )
        self._index = DuplicateIndex()
        for number, row in enumerate(self._snapshot, start=2):
            record = dict(zip(CATALOGUE_COLUMNS, row))
            self._index.add(
                _text(record["Authors"]), _text(record["Link"]), _text(record["Title"]), f"row {number}"
            )

    @property
    def last_row(self) -> int:
        return self._last_row

    def plan(self, records: list[ParsedRecord], force: bool = False) -> MergePlan:
        """
        Sort records into those to write and duplicates. Records to write join
        the index, so later inputs are checked against them as well.
        """
        plan = MergePlan()
        for record in records:
            duplicate = self._index.find(record)
            if duplicate and not force:
                plan.skipped.append(duplicate)
                continue
            if duplicate:
                plan.forced.append(duplicate)
            plan.to_write.append(record)
            metadata = record.metadata
            self._index.add(metadata.Authors, metadata.Link, metadata.Title, record.label)
        return plan

    def write(self, plan: MergePlan) -> MergeReport:
        """Append the planned rows, then save, validate and replace the catalogue."""
        if not plan.to_write:
            raise CatalogueError("Nothing to write.")

        template = self._last_row
        first_row = self._last_row + 1
        expected: dict[int, list] = {}
        for offset, record in enumerate(plan.to_write):
            row = first_row + offset
            values = row_values(self._next_no + offset, record.metadata)
            self._write_row(row, values, template)
            expected[row] = values
            logger.info("Row %s: %s (%s)", row, record.metadata.Title, record.label)

        last_row = first_row + len(plan.to_write) - 1
        self._table.ref = _table_ref(last_row)
        if self._table.autoFilter is not None:
            self._table.autoFilter.ref = self._table.ref

        self._save_validated(expected, last_row)
        return MergeReport(first_row, last_row)

    def _write_row(self, row: int, values: list, template: int) -> None:
        for column, value in enumerate(values, start=1):
            cell = self._sheet.cell(row=row, column=column, value=value)
            if isinstance(value, str):
                # Keep text such as "=…" a string instead of a formula.
                cell.data_type = "s"
            source = self._sheet.cell(row=template, column=column)
            if template > 1 and source.has_style:
                cell._style = copy(source._style)
        if template > 1:
            self._sheet.row_dimensions[row].height = self._sheet.row_dimensions[template].height

    def _save_validated(self, expected: dict[int, list], last_row: int) -> None:
        handle, tmp_name = tempfile.mkstemp(
            dir=self._path.parent, prefix=f".{self._path.stem}.", suffix=self._path.suffix
        )
        os.close(handle)
        tmp_path = Path(tmp_name)
        try:
            self._workbook.save(tmp_path)
            validate_merge(tmp_path, self._snapshot, expected, last_row)
            os.replace(tmp_path, self._path)
        except BaseException:
            tmp_path.unlink(missing_ok=True)
            raise


def validate_merge(
    path: Path, snapshot: list[tuple], expected: dict[int, list], last_row: int
) -> None:
    """
    Check a merged catalogue: header and table range are right, existing rows
    are unchanged and every new cell sits in its column.
    """
    workbook = _open_workbook(path)
    sheet = _catalogue_sheet(workbook, path)
    _check_header(sheet, path)

    table = _catalogue_table(sheet, path)
    if table.ref != _table_ref(last_row):
        raise CatalogueError(
            f"Validation failed: table range is {table.ref}, expected {_table_ref(last_row)}."
        )

    for number, (before, after) in enumerate(zip(snapshot, _values(sheet, 2, len(snapshot) + 1)), start=2):
        if before != after:
            _raise_mismatch(number, before, after, "existing row changed")

    first_new = len(snapshot) + 2
    for number, actual in enumerate(_values(sheet, first_new, last_row), start=first_new):
        if list(actual) != expected[number]:
            _raise_mismatch(number, expected[number], actual, "merged row differs")


def _raise_mismatch(number: int, expected, actual, what: str) -> None:
    for column, (want, got) in enumerate(zip(expected, actual), start=1):
        if want != got:
            raise CatalogueError(
                f"Validation failed ({what}) at {get_column_letter(column)}{number} "
                f"[{CATALOGUE_COLUMNS[column - 1]}]: expected {want!r}, found {got!r}."
            )
    raise CatalogueError(f"Validation failed ({what}) in row {number}.")


def _open_workbook(path: Path) -> Workbook:
    try:
        return load_workbook(path)
    except (InvalidFileException, BadZipFile, KeyError, OSError) as exc:
        raise CatalogueError(f"Cannot open catalogue {path}: {exc}") from exc


def _catalogue_sheet(workbook: Workbook, path: Path) -> Worksheet:
    if CATALOGUE_SHEET_NAME not in workbook.sheetnames:
        raise CatalogueError(f"{path} has no sheet '{CATALOGUE_SHEET_NAME}'.")
    return workbook[CATALOGUE_SHEET_NAME]


def _check_header(sheet: Worksheet, path: Path) -> None:
    header = [cell.value for cell in sheet[1][: len(CATALOGUE_COLUMNS)]]
    if header != CATALOGUE_COLUMNS:
        raise CatalogueError(
            f"{path} does not carry the expected header.\n"
            f"Expected: {', '.join(CATALOGUE_COLUMNS)}\n"
            f"Found:    {', '.join(str(value) for value in header)}"
        )


def _catalogue_table(sheet: Worksheet, path: Path) -> Table:
    for table in sheet.tables.values():
        min_col, min_row, max_col, _ = range_boundaries(table.ref)
        if (min_col, min_row, max_col) == (1, 1, len(CATALOGUE_COLUMNS)):
            return table
    raise CatalogueError(
        f"{path} has no table spanning columns A:{get_column_letter(len(CATALOGUE_COLUMNS))} "
        f"from row 1 on sheet '{CATALOGUE_SHEET_NAME}'."
    )


def _table_ref(last_row: int) -> str:
    return f"A1:{get_column_letter(len(CATALOGUE_COLUMNS))}{last_row}"


def _values(sheet: Worksheet, first_row: int, last_row: int) -> list[tuple]:
    if last_row < first_row:
        return []
    return list(
        sheet.iter_rows(
            min_row=first_row,
            max_row=last_row,
            max_col=len(CATALOGUE_COLUMNS),
            values_only=True,
        )
    )


def _text(value) -> str:
    return "" if value is None else str(value)
